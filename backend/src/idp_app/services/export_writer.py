"""Disk-spooled relational exports; memory is bounded to one source and one row.

Preflight discovers data-dependent columns before any worksheet rows are emitted.
Spools are private temporary files, never user-controlled paths.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
import sqlite3
import zipfile
from collections.abc import Iterable
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from idp_app.services.document_models import ExtractionRunRecord
from idp_app.services.exports import ExportTable, data_dictionary_rows
from idp_app.services.schema_models import SchemaRecord

WRITER_VERSION = "2"
MAX_ROWS = 1_048_576
MAX_COLUMNS = 16_384
MAX_CELL = 32_767


def _name(label: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", label).strip("'")[:31] or "Records"
    result = base
    number = 1
    while result.casefold() in used:
        number += 1
        result = f"{base[: 31 - len(str(number)) - 1]}_{number}"
    used.add(result.casefold())
    return result


def _value(value: Any) -> Any:
    return json.dumps(value, ensure_ascii=False) if isinstance(value, dict | list) else value


def _cells(sheet: Any, values: list[Any], *, header: bool = False) -> list[Any]:
    result = []
    for value in values:
        value = _value(value)
        if isinstance(value, str) and len(value) > MAX_CELL:
            raise ValueError("Excel cells support at most 32,767 characters; export CSV instead.")
        cell = WriteOnlyCell(sheet, value=value)
        if isinstance(value, str):
            cell.data_type = "s"  # Formula-looking evidence and leading zero IDs remain text.
        if header:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="202420")
        result.append(cell)
    return result


def write_export(
    sources: Iterable[tuple[ExtractionRunRecord, SchemaRecord, list[ExportTable]]],
    destination: Path,
    format: str,
    *,
    max_disk_bytes: int = 2_000_000_000,
    row_limit: int = MAX_ROWS,
) -> bool:
    """Return whether the output spans schema versions. Caller owns destination cleanup."""
    if format not in {"xlsx", "csv"} or not 2 <= row_limit <= MAX_ROWS:
        raise ValueError("Invalid export format or row limit")
    with TemporaryDirectory(prefix="idp-export-", dir=destination.parent) as directory:
        root = Path(directory)
        connection = sqlite3.connect(root / "rows.sqlite")
        connection.execute("CREATE TABLE rows (table_id TEXT, payload TEXT)")
        connection.execute("CREATE INDEX rows_table ON rows(table_id)")
        groups: dict[tuple[str, int], dict[str, Any]] = {}
        written = 0
        try:
            for run, schema, tables in sources:
                group = groups.setdefault(
                    (run.schema_id, run.schema_version),
                    {
                        "schema": schema,
                        "tables": {},
                    },
                )
                for table in tables:
                    identity = hashlib.sha256(
                        json.dumps(
                            [
                                run.schema_id,
                                run.schema_version,
                                table.schema_path,
                            ]
                        ).encode()
                    ).hexdigest()
                    meta = group["tables"].setdefault(
                        table.schema_path,
                        {
                            "id": identity,
                            "name": table.name,
                            "columns": [],
                            "count": 0,
                        },
                    )
                    for column in table.columns:
                        if column not in meta["columns"]:
                            meta["columns"].append(column)
                    if format == "xlsx" and len(meta["columns"]) > MAX_COLUMNS:
                        raise ValueError(
                            "Excel supports at most 16,384 columns; export CSV instead."
                        )
                    for row in table.rows:
                        payload = json.dumps(row, ensure_ascii=False)
                        written += len(payload.encode())
                        if written > max_disk_bytes // 3:
                            raise ValueError(
                                "Export exceeds the temporary disk budget; select fewer results."
                            )
                        connection.execute("INSERT INTO rows VALUES (?, ?)", (identity, payload))
                        meta["count"] += 1
                connection.commit()
            multiple = len(groups) > 1
            outputs: list[tuple[str, Path]] = []
            for index, ((schema_id, version), group) in enumerate(groups.items()):
                suffix = "xlsx" if format == "xlsx" else "zip"
                output = root / f"group-{index}.{suffix}"
                # Schema identifiers are labels, never filesystem paths.
                label = re.sub(r"[^A-Za-z0-9_.-]", "_", schema_id)[:100]
                label = f"{label}_v{version}_{index}.{suffix}"
                _write_group(connection, group, output, format, row_limit)
                outputs.append((label, output))
                if sum(p.stat().st_size for p in root.iterdir()) > max_disk_bytes:
                    raise ValueError(
                        "Export exceeds the temporary disk budget; select fewer results."
                    )
            if len(outputs) == 1:
                shutil.copyfile(outputs[0][1], destination)
            else:
                with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
                    for label, output in outputs:
                        if format == "xlsx":
                            archive.write(output, label)
                        else:
                            with zipfile.ZipFile(output) as inner:
                                for entry in inner.namelist():
                                    with (
                                        inner.open(entry) as src,
                                        archive.open(f"{label[:-4]}/{entry}", "w") as dst,
                                    ):
                                        shutil.copyfileobj(src, dst, 1024 * 1024)
            return multiple
        finally:
            connection.close()


def _write_group(
    connection: sqlite3.Connection, group: dict[str, Any], output: Path, format: str, row_limit: int
) -> None:
    used = {"data_dictionary", "export_manifest"}
    dictionary = [
        ["Schema path", "Field name", "Type", "Description"],
        *[list(row) for row in data_dictionary_rows(group["schema"])],
    ]
    manifest: list[list[Any]] = [["Sheet", "Schema path", "Rows"]]
    workbook = Workbook(write_only=True) if format == "xlsx" else None
    archive = zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) if format == "csv" else None

    def sheet_rows(name: str, columns: list[str], rows: Iterable[list[Any]], count: int) -> None:
        if workbook is not None:
            sheet = workbook.create_sheet(name)
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for index, column in enumerate(columns, 1):
                sheet.column_dimensions[get_column_letter(index)].width = min(
                    38, max(12, len(column) + 2)
                )
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"
            sheet.append(_cells(sheet, columns, header=True))
            try:
                for row in rows:
                    sheet.append(_cells(sheet, row))
            finally:
                sheet.close()
        else:
            assert archive is not None
            with (
                archive.open(f"{name}.csv", "w") as stream,
                io.TextIOWrapper(stream, encoding="utf-8", newline="") as text,
            ):
                writer = csv.writer(text)
                writer.writerow(columns)
                for row in rows:
                    # Quote spreadsheet formulas as data when opened in Excel.
                    writer.writerow(
                        [
                            "'" + value
                            if isinstance(value := _value(v), str)
                            and value.startswith(("=", "+", "-", "@"))
                            else value
                            for v in row
                        ]
                    )

    try:
        for meta in group["tables"].values():
            columns = meta["columns"]
            chunk_size = row_limit - 1 if workbook is not None else max(1, meta["count"])
            cursor = connection.execute(
                "SELECT payload FROM rows WHERE table_id = ? ORDER BY rowid", (meta["id"],)
            )
            for offset in range(0, max(1, meta["count"]), chunk_size):
                count = min(chunk_size, meta["count"] - offset)
                name = _name(meta["name"], used)

                def rows(
                    count: int = count,
                    cursor: sqlite3.Cursor = cursor,
                    columns: list[str] = columns,
                ) -> Iterable[list[Any]]:
                    for _ in range(count):
                        row = json.loads(cursor.fetchone()[0])
                        yield [row.get(column) for column in columns]

                sheet_rows(name, columns, rows(), count)
                manifest.append(
                    [name, next(k for k, v in group["tables"].items() if v is meta), count]
                )
        sheet_rows("Data_dictionary", dictionary[0], dictionary[1:], len(dictionary) - 1)
        sheet_rows("Export_manifest", manifest[0], manifest[1:], len(manifest) - 1)
        if workbook is not None:
            workbook.save(output)
    finally:
        if archive is not None:
            archive.close()
        if workbook is not None:
            workbook.close()
