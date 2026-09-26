import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_generic_exports import _nested_schema, _run

from idp_app.services.export_writer import write_export
from idp_app.services.exports import ExportTable


def test_spool_headers_splits_collisions_and_literal_cells(tmp_path: Path) -> None:
    def sources():
        for i in range(5):
            yield (
                _run(),
                _nested_schema(),
                [
                    ExportTable(
                        "Same",
                        "$.a",
                        ["id", "value", *(["late"] if i == 4 else [])],
                        [{"id": "001", "value": "=1+1", "late": i}],
                    ),
                    ExportTable("Same", "$.b", ["value"], [{"value": -4}]),
                ],
            )

    output = tmp_path / "result.xlsx"
    assert not write_export(sources(), output, "xlsx", row_limit=3)
    book = load_workbook(output)
    assert len(book.sheetnames) == 8
    assert list(book["Same"].values)[0] == ("id", "value", "late")
    assert book["Same"]["B2"].data_type == "s"
    assert book["Same"]["A2"].value == "001"
    assert book["Same"].auto_filter.ref == "A1:C3"
    assert sum(row[2] for row in list(book["Export_manifest"].values)[1:]) == 10
    book.close()


def test_oversize_text_errors_and_lossless_csv(tmp_path: Path) -> None:
    tables = [ExportTable("Data", "$", ["text"], [{"text": "x" * 32768}])]
    output = tmp_path / "result"
    with pytest.raises(ValueError, match="32,767"):
        write_export([(_run(), _nested_schema(), tables)], output, "xlsx")
    write_export([(_run(), _nested_schema(), tables)], output, "csv")
    with zipfile.ZipFile(output) as archive:
        assert "x" * 32768 in archive.read("Data.csv").decode()


def test_disk_budget_is_enforced(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="disk budget"):
        write_export(
            [
                (
                    _run(),
                    _nested_schema(),
                    [ExportTable("Data", "$", ["text"], [{"text": "x" * 1000}])],
                )
            ],
            tmp_path / "result",
            "csv",
            max_disk_bytes=100,
        )
